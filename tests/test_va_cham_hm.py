"""Va chạm (va_cham.py) + kiểm nhiều hạng mục (hm_pages.py, check từng hạng mục) — offline."""
import json
from pathlib import Path

import pymupdf
import pytest

from pccc_checker.cli import main
from pccc_checker.models import BuildingInfo, DrawingExtraction, PageExtraction
from pccc_checker.pipeline import hm_pages, va_cham
from pccc_checker.project_info import CONFIRMED, FieldValue, ProjectInfo, load_spec

HM = ["Nhà văn phòng", "Nhà xưởng 1", "Nhà xưởng 2", "Nhà xưởng 3", "Nhà xe 01"]


@pytest.mark.parametrize("line,want", [
    ("NHÀ XƯỞNG 01 VÀ NHÀ VĂN PHÒNG", {"Nhà văn phòng", "Nhà xưởng 1"}),
    ("NHÀ XƯỞNG CHO THUÊ 02, 03:  2000 M2", {"Nhà xưởng 2", "Nhà xưởng 3"}),
    ("MẶT BẰNG BÁO CHÁY TẦNG 1 NHÀ XƯỞNG 03", {"Nhà xưởng 3"}),
    ("CHI TIẾT NHÀ XE 01", {"Nhà xe 01"}),
    ("NHÀ XƯỞNG CHO THUÊ", set()),
    ("NHÀ XƯỞNG 12", set()),
])
def test_line_owners(line, want):
    assert hm_pages.line_owners(line, HM) == want


def test_page_owners_ke_thua_va_chung():
    lines = {1: ["BÌA"], 2: ["MẶT BẰNG TỔNG THỂ", "NHÀ XƯỞNG 01"], 3: ["NHÀ XƯỞNG 01 VÀ NHÀ VĂN PHÒNG"],
             4: ["MẶT ĐỨNG"], 5: ["NHÀ XƯỞNG CHO THUÊ 02, 03"], 6: ["MẶT CẮT"],
             7: ["NHÀ VĂN PHÒNG", "NHÀ XƯỞNG 2", "NHÀ XE 01"]}
    ow = hm_pages.page_owners(lines, HM, cats={2: "site"}, titles={2: "MẶT BẰNG TỔNG THỂ"})
    assert ow[1] == ["chung"] and ow[2] == ["chung"]
    assert ow[3] == ow[4] == ["Nhà văn phòng", "Nhà xưởng 1"]
    assert ow[5] == ow[6] == ["Nhà xưởng 2", "Nhà xưởng 3"]
    assert ow[7] == ["chung"]                                # 3 loại hạng mục trên 1 trang -> tổng thể
    assert hm_pages.by_hm(ow)["Nhà xưởng 2"] == [5, 6]


def _fv(v):
    return FieldValue(v, None, 1, "x", "user", CONFIRMED)


def _info() -> ProjectInfo:
    a = {"nhom_nha": _fv(["F5.1"]), "chieu_cao_pccc": _fv(6.5), "bac_chiu_lua": _fv("IV"),
         "cap_nguy_hiem_chay_ket_cau": _fv("S0"), "so_tang_noi": _fv(2), "hang_chay_no": _fv("C")}
    b = {"chieu_cao_pccc": _fv(5.97)}                        # thiếu nhóm nhà, bậc chịu lửa…
    info = ProjectInfo(hang_muc={"Nhà xưởng 2": a, "Nhà văn phòng": b}, hang_muc_chinh="Nhà xưởng 2")
    for k, v in a.items():
        info.fire_inputs[k] = v
    return info


def test_for_hang_muc_lay_thong_so_rieng():
    spec = load_spec()
    vp = _info().for_hang_muc("Nhà văn phòng", spec)
    assert vp.get("chieu_cao_pccc").value == 5.97 and vp.get("nhom_nha").value is None
    assert vp.to_building(spec)["height_pccc_m"] == 5.97 and "nhom_nha" in vp.problems(spec)


def test_check_chay_rieng_tung_hang_muc(tmp_path, capsys):
    d = DrawingExtraction("x.pdf", BuildingInfo(), [PageExtraction(p, 1000, 700) for p in (1, 2, 3)])
    (tmp_path / "extraction.json").write_text(json.dumps(d.to_dict()), encoding="utf-8")
    (tmp_path / "project_info.json").write_text(json.dumps(_info().to_dict(), ensure_ascii=False), encoding="utf-8")
    (tmp_path / "hm_pages.json").write_text(json.dumps(
        {"trang": {"1": ["chung"], "2": ["Nhà xưởng 2"], "3": ["Nhà văn phòng"]}}, ensure_ascii=False), encoding="utf-8")
    main(["check", str(tmp_path / "extraction.json")])
    fs = json.loads((tmp_path / "findings.json").read_text(encoding="utf-8"))
    by = {}
    for f in fs:
        by.setdefault(f["hang_muc"], []).append(f)
    out = capsys.readouterr().out
    assert all(f"  {n}: 1 trang" in out for n in ("chung", "Nhà xưởng 2", "Nhà văn phòng"))
    # văn phòng thiếu nhóm nhà -> luật không kiểm được; xưởng 2 đủ thông số -> không có
    assert by["Nhà văn phòng"] and all(f["status"] == "cannot_check" or f["rule_id"] == "BUILDING-INFO"
                                       for f in by["Nhà văn phòng"])
    assert not any(f["status"] == "cannot_check" for f in by.get("Nhà xưởng 2", []))
    assert "Hạng mục" in (tmp_path / "report.md").read_text(encoding="utf-8")
    main(["check", str(tmp_path / "extraction.json"), "--mot-hm"])
    assert {f["hang_muc"] for f in json.loads((tmp_path / "findings.json").read_text(encoding="utf-8"))} <= {""}
    assert ": 1 trang" not in capsys.readouterr().out          # cách cũ: không chia theo hạng mục


# ------------------------------------------------------------------ va chạm
@pytest.mark.parametrize("v,toks", [
    (2000, {"2000", "2.000", "2 000"}), (11.425, {"11.425", "11,425"}),
    ("3000 m², 36.510 m³", {"3000", "36510", "36.510", "36 510"}),
])
def test_tokens(v, toks):
    assert toks <= set(va_cham._tokens(v))


def _png(w=900, h=300) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), False)
    pix.set_rect(pix.irect, (230, 230, 230))
    return pix.tobytes("png")


def _pdf(tmp_path: Path) -> Path:
    doc = pymupdf.open()
    p = doc.new_page(width=842, height=595)                                    # 1: chữ vector
    p.insert_text((100, 100), "DIỆN TÍCH XÂY DỰNG: 2000 m2", fontname="tiro", fontsize=12)
    p.insert_text((100, 300), "THUYẾT MINH: DIỆN TÍCH 3000 m2", fontname="tiro", fontsize=12)
    p = doc.new_page(width=842, height=595)                                    # 2: bảng ảnh, trang xoay 270
    p.insert_image(pymupdf.Rect(100, 100, 700, 300), stream=_png())
    p.set_rotation(270)
    f = tmp_path / "t.pdf"
    doc.save(f)
    return f


def _out(tmp_path: Path) -> Path:
    out = tmp_path / "out"
    out.mkdir()
    (out / "nghi_ngo.json").write_text(json.dumps({"mau_thuan_tren_ban_ve": [
        {"hang_muc": "Nhà xưởng 1", "truong": "dien_tich_xay_dung", "trang": [1],
         "gia_tri": {"bảng tr.1": 2000, "thuyết minh tr.1": "3000 m2"}},
        {"hang_muc": "Nhà xưởng 1", "truong": "so_tang", "gia_tri": {"bảng A tr.2": 2, "bảng B tr.2": 1}},
        {"hang_muc": "Nhà xưởng 1", "truong": "chieu_cao_cong_trinh", "gia_tri": {"tr.1": 11.425, "tr.2": 11.43}},
    ]}, ensure_ascii=False), encoding="utf-8")
    info = {"hang_muc": {"Nhà xưởng 1": {"so_tang_noi": {"status": "conflict", "candidates": [
                {"value": 2, "page": 2, "method": "gemini"}, {"value": 1, "page": 2, "method": "gemini"}]},
            "bac_chiu_lua": {"status": "conflict", "candidates": [
                {"value": "IV", "page": 1, "method": "vector"}, {"value": "II", "page": 2, "method": "gemini"}]}}},
            "khu_dat": {}}
    (out / "project_info.json").write_text(json.dumps(info, ensure_ascii=False), encoding="utf-8")
    return out


def test_collect_gop_nguon_va_lam_tron(tmp_path):
    items = va_cham.collect(_out(tmp_path))
    assert [(i["truong"], i["loai"]) for i in items] == [
        ("dien_tich_xay_dung", "ban_ve"), ("so_tang", "ban_ve"), ("chieu_cao_cong_trinh", "ban_ve"),
        ("bac_chiu_lua", "doc_lech")]                       # so_tang_noi trùng so_tang của nghi_ngo -> bỏ
    assert items[2]["bo"].startswith("làm tròn") and items[0]["gia_tri"][1]["vt"] == "VC-001.2"


def test_locate_annotate_trang_xoay(tmp_path):
    pdf, out = _pdf(tmp_path), _out(tmp_path)
    items = va_cham.collect(out)
    va_cham.locate(items, pdf)
    a, b = items[0]["gia_tri"]
    assert a["vi_tri"]["cach"] == b["vi_tri"]["cach"] == "vector"
    assert a["vi_tri"]["bbox"][1] < 0.2 < 0.45 < b["vi_tri"]["bbox"][1]       # 2000 ở y=100, 3000 ở y=300
    s1, s2 = items[1]["gia_tri"]
    assert s1["vi_tri"]["cach"] == "vung_anh" and s1["vi_tri"]["anh"] == 1
    assert any(x.startswith("VC-002.1") for x in va_cham.can_claude(items))
    log = va_cham.apply_claude(items, [{"vt": "VC-002.2", "anh": 1, "bbox_trong_anh": [0, 0, 1, 0.5]},
                                       {"id": "VC-004", "bo": True, "ly_do": "Gemini"}], pdf)
    assert len(log) == 2 and items[3]["bo"] == "Gemini"
    box = s2["vi_tri"]["bbox"]
    # ảnh 600x200 trên trang chưa xoay -> trang hiển thị 595x842: ảnh dọc 200x600; nửa trên ảnh cắt = 200x300
    assert box[2] - box[0] == pytest.approx(200 / 595, abs=0.01) and box[3] - box[1] == pytest.approx(300 / 842, abs=0.01)
    n = va_cham.annotate(pdf, items, tmp_path / "vc.pdf")
    assert n == 4                                          # 2 va chạm × 2 vị trí; làm tròn + doc_lech đã bỏ
    doc = pymupdf.open(tmp_path / "vc.pdf")
    page = doc[1]
    r = next(x.rect for x in page.annots() if x.type[1] == "Square" and "VC-002.2" in x.info["title"])
    vis = r * page.rotation_matrix                          # quay lại toạ độ hiển thị
    W, H = page.rect.width, page.rect.height
    assert vis.x0 / W == pytest.approx(box[0], abs=0.02) and vis.y1 / H == pytest.approx(box[3], abs=0.02)
    assert "VC-001" in va_cham.to_md(items)


def test_cli_va_cham(tmp_path, capsys):
    pdf, out = _pdf(tmp_path), _out(tmp_path)
    main(["va-cham", str(out), "--pdf", str(pdf), "--annotate"])
    assert (out / "va_cham.json").exists() and (out / "va_cham.md").exists() and (out / "t_va_cham.pdf").exists()
    from docx import Document
    doc = Document(out / "Note phản hồi.docx")
    txt = " | ".join(p.text for p in doc.paragraphs)
    assert "Note phản hồi" in txt and "Tr.1: VC-001.1, VC-001.2" in txt
    t = doc.tables[0]
    assert [c.text for c in t.rows[0].cells][:3] == ["Trang", "Mã note", "Loại"]
    pages = [r.cells[0].text for r in t.rows[1:]]
    assert pages == sorted(pages, key=int) and len(pages) == 6      # 3 va chạm (bỏ làm tròn) × 2 vị trí, xếp theo trang
    assert "Claude cần chỉ vị trí" in capsys.readouterr().out


def test_clear_env_cookies(tmp_path, monkeypatch):
    from pccc_checker.extractors.gemini_web import clear_env_cookies
    env = tmp_path / ".env"
    env.write_text("GEMINI_API_KEY=x\nGEMINI_COOKIES=[1]\nGEMINI_1PSID=a\n", encoding="utf-8")
    monkeypatch.setenv("GEMINI_COOKIES", "[1]")
    clear_env_cookies(env)
    assert env.read_text(encoding="utf-8") == "GEMINI_API_KEY=x\n"
    import os
    assert "GEMINI_COOKIES" not in os.environ
