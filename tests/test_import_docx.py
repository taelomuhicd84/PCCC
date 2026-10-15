"""scripts/knowled_import_pdf.py nạp .docx: đánh số tự động theo kiểu Heading, tách mục, bảng có tiêu đề thành module riêng."""
import importlib.util
from pathlib import Path

import docx

SP = Path(__file__).resolve().parents[1] / "scripts" / "knowled_import_pdf.py"
spec = importlib.util.spec_from_file_location("knowled_import_pdf", SP)
imp = importlib.util.module_from_spec(spec)
spec.loader.exec_module(imp)


def _mau(tmp_path):
    d = docx.Document()
    d.add_paragraph("MỤC LỤC")
    d.add_paragraph("PHỤ LỤC A Quy định về trang bị báo cháy ………..9")
    d.add_heading("QUY ĐỊNH CHUNG", 1)
    d.add_heading("Phạm vi điều chỉnh", 2)
    d.add_heading("Quy chuẩn này quy định về trang bị phương tiện PCCC cho nhà và công trình, kể cả công trình đang thi công " * 2, 3)
    d.add_heading("Tầng nổi", 2)
    d.add_paragraph("Bảng 1 - Giới hạn chịu lửa")
    t = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("Bậc", "R"), ("I", "R 120"), ("IV", "R 15")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    d.add_paragraph("Quy định về trang bị báo cháy")
    d.add_paragraph("Nội dung phụ lục A")
    f = tmp_path / "mau.docx"
    d.save(f)
    return f


def test_danh_so_heading_va_tach_muc(tmp_path):
    lines, bang = imp.docx_parts(_mau(tmp_path))
    assert any(ln.lstrip("\x01").startswith("1.1 Phạm vi") for ln in lines)
    assert any(ln.lstrip("\x01").startswith("1.1.1 Quy chuẩn này") for ln in lines)       # điều khoản dài vẫn là tiêu đề
    assert any(ln.lstrip("\x01").startswith("1.2 Tầng nổi") for ln in lines)
    muc = {n: (t, c) for n, t, c in imp.tach_muc(lines, "qcx", 3)}
    assert {"1", "1.1", "1.1.1", "1.2"} <= set(muc)
    assert "[Bảng → qcx-bang-1]" in muc["1.2"][1]
    assert "pl-a" in muc and "Nội dung phụ lục A" in muc["pl-a"][1]                    # tiêu đề phụ lục lấy từ mục lục
    assert bang["1"]["rows"][2] == ["IV", "R 15"]
