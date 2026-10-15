"""Điền kết quả đối chiếu vào Word mẫu datamau: Nội dung thiết kế + Kết luận đúng hàng (R### theo src_map)."""
import json

from docx import Document

from pccc_checker.output import bang_doi_chieu as b


def test_dien_vao_dung_hang_va_ket_luan(tmp_path):
    sm = json.load(open(b.CHECKLISTS / "khach-san" / "src_map.json", encoding="utf-8"))
    ids = [i for i in sm if i.startswith("KS1-QM-")][1:4]
    fs = [dict(hm="HM01", rule=ids[0], tt="D", trang=7, ly_do="Cao PCCC 20,7 m", evid="CHIỀU CAO PCCC 20,7"),
          dict(hm="HM01", rule=ids[1], tt="KD", trang=9, ly_do="Thiếu 1 lối", evid=""),
          dict(hm="HM01", rule=ids[2], tt="NA", trang=None, ly_do="Không có hầm", evid="")]
    doc = Document(b._mau_file("B7.1"))
    st = b.dien(doc, sm, fs, "KS1-")
    assert st["dien"] == 3
    rows = b.hang_theo_R(doc)
    kl = [rows[sm[i]["row"]][5].text for i in ids]
    assert kl == ["Đạt", "Không đạt (tr.9)", "Không áp dụng"]
    assert "tr.7" in rows[sm[ids[0]]["row"]][2].text and "20,7" in rows[sm[ids[0]]["row"]][2].text
    assert b.ket_luan(fs[1])[1] == b.RED


def test_font_tieng_viet_khong_bi_doi_dau(tmp_path):
    doc = Document(b._mau_file("B7.1"))
    c = b.hang_theo_R(doc)["R003"]
    b._ghi(c[2], "Khoảng cách tới nhà dân, tường cao tới mái")
    rpr = c[2].paragraphs[0].runs[-1]._element.rPr
    assert rpr.rFonts.get(b.qn("w:ascii")) == "Times New Roman" and rpr.rFonts.get(b.qn("w:hAnsi")) == "Times New Roman"


def test_dau_bang_font_tuong_minh():
    doc = Document(b._mau_file("B7.1"))
    b.dien_dau_bang(doc, {"ten": "HOMESTAY NGUYỄN KHOÁI (NHÀ NGHỈ)", "chu_dau_tu": "Ông Nguyễn Thế Vương"})
    ps = [p for p in doc.paragraphs[:20] if p.text.startswith(("1. Tên", "3. Chủ"))]
    assert len(ps) == 2
    for p in ps:
        assert all(r._element.rPr.rFonts.get(b.qn("w:ascii")) == "Times New Roman" for r in p.runs)


def test_chuan_font_va_gac_run_thieu_font():
    """Gốc lỗi: style Normal của mẫu = .VnTime; run không ghi font -> méo dấu. chuan_font phải vá hết."""
    doc = Document(b._mau_file("B7.1"))
    doc.add_paragraph().add_run("Khoảng cách tới nhà dân")                 # run không font (như p.text = ...)
    truoc = b.font_loi(doc)                                                # mẫu có sẵn vài run ký tự ASCII không font
    assert truoc >= 1
    assert b.chuan_font(doc) == truoc and b.font_loi(doc) == 0


def test_sach_tu_noi_bo():
    assert "gói" not in b._sach("Thiếu bảng trong gói quy định") and "KS1-TN-012" not in b._sach("phụ thuộc mục KS1-TN-012")
