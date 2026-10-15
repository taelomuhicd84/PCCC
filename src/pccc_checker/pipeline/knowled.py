"""knowled — kho kiến thức để Claude nạp & so sánh (skill `knowled`).

Một nơi tra duy nhất, nạp theo nhu cầu (không nạp cả kho):
  knowledge/<nhóm>/<id>.md      mục kiến thức do người dùng/Claude nạp (front matter: id, tieu_de, nhom, tu_khoa, trang_thai, nguon)
  knowledge/_nguon.yaml         file sẵn có coi như kho (định nghĩa QCVN, bài học…) — tách theo tiêu đề, không sao chép
Trạng thái: `verified` (đối chiếu nguyên văn / người dùng xác nhận) | `cho_duyet` (mới nạp). Kiến thức chờ duyệt chỉ được dùng
để gợi ý (XM), KHÔNG dùng để kết luận KD.
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date
from pathlib import Path

import yaml

from ..config import PROJECT_ROOT

KHO = PROJECT_ROOT / "knowledge"
NHOM = ("cong-nang", "quy-chuan", "huong-dan", "kinh-nghiem", "thieu")
FM = re.compile(r"\A---\s*\n(.*?)\n---\s*\n?", re.S)


def bo_dau(s: str) -> str:
    s = unicodedata.normalize("NFD", (s or "").replace("đ", "d").replace("Đ", "D"))
    return "".join(c for c in s if unicodedata.category(c) != "Mn").lower()


def _fm(text: str) -> tuple[dict, str]:
    m = FM.match(text)
    return (yaml.safe_load(m.group(1)) or {}, text[m.end():]) if m else ({}, text)


def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", bo_dau(s)).strip("-")[:48]


def _nguon_ngoai() -> list[Path]:
    f = KHO / "_nguon.yaml"
    return [PROJECT_ROOT / p for p in (yaml.safe_load(f.read_text(encoding="utf-8")) or [])] if f.exists() else []


def muc_kho() -> list[dict]:
    """Mọi mục: {id, tieu_de, nhom, tu_khoa, trang_thai, nguon, noi_dung, file}."""
    out = []
    for f in sorted(KHO.rglob("*.md")):
        if f.name.startswith("_"):
            continue
        fm, body = _fm(f.read_text(encoding="utf-8-sig"))
        out.append(dict(id=fm.get("id") or f.stem, tieu_de=fm.get("tieu_de") or f.stem, nhom=fm.get("nhom") or f.parent.name,
                        tu_khoa=fm.get("tu_khoa") or "", trang_thai=fm.get("trang_thai") or "cho_duyet", nguon=fm.get("nguon") or "",
                        noi_dung=body.strip(), file=f))
    for f in _nguon_ngoai():                                   # file sẵn có: tách theo tiêu đề markdown
        if not f.exists():
            continue
        cur, buf, n = f.stem, [], 0
        for ln in f.read_text(encoding="utf-8-sig").splitlines() + ["# ~"]:
            if ln.startswith("#"):
                if "".join(buf).strip():
                    n += 1
                    out.append(dict(id=f"{f.stem}#{_slug(cur) or n}", tieu_de=cur, nhom="quy-chuan" if "qcvn" in str(f).lower() else "kinh-nghiem",
                                    tu_khoa="", trang_thai="verified", nguon=str(f.relative_to(PROJECT_ROOT)), noi_dung="\n".join(buf).strip(), file=f))
                cur, buf = ln.lstrip("# ").strip(), []
            else:
                buf.append(ln)
    return out


def tim(q: str, top: int = 5, nhom: str | None = None) -> list[tuple[float, dict, str]]:
    """Tra từ khoá (không phân biệt dấu/hoa thường). -> [(điểm, mục, đoạn trích)]."""
    ws = [w for w in re.findall(r"[\w.]+", bo_dau(q)) if len(w) > 1]
    kq = []
    for m in muc_kho():
        if nhom and m["nhom"] != nhom:
            continue
        tk, td, nd = bo_dau(str(m["tu_khoa"])), bo_dau(m["tieu_de"]), bo_dau(m["noi_dung"])
        sc = sum(5 * (w in tk) + 3 * (w in td) + min(nd.count(w), 4) for w in ws)
        sc += 4 * sum(w in nd for w in ws) // max(len(ws), 1) if ws and all(w in nd for w in ws) else 0
        if sc <= 0:
            continue
        i = min((nd.find(w) for w in ws if w in nd), default=0)
        kq.append((sc, m, re.sub(r"\s+", " ", m["noi_dung"][max(i - 60, 0): i + 180]).strip()))
    return sorted(kq, key=lambda x: -x[0])[:top]


def lay(id_: str) -> dict | None:
    return next((m for m in muc_kho() if m["id"] == id_), None)


def docx_sang_md(f: Path) -> str:
    import docx
    d = docx.Document(f)
    out = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        out.append("")
        for r in t.rows:
            out.append("| " + " | ".join(re.sub(r"\s+", " ", c.text).strip() for c in r.cells) + " |")
    return "\n".join(out)


def them(nguon_file: Path, id_: str | None, tieu_de: str, nhom: str, tu_khoa: str, nguon: str, duyet: bool = False) -> Path:
    """Nạp 1 file (.md/.txt/.docx) vào knowledge/<nhóm>/<id>.md (mặc định chờ duyệt)."""
    if nhom not in NHOM:
        raise SystemExit(f"nhóm phải thuộc {NHOM}")
    body = docx_sang_md(nguon_file) if nguon_file.suffix.lower() == ".docx" else nguon_file.read_text(encoding="utf-8-sig")
    id_ = id_ or _slug(tieu_de)
    p = KHO / nhom / f"{id_}.md"
    p.parent.mkdir(parents=True, exist_ok=True)
    fm = dict(id=id_, tieu_de=tieu_de, nhom=nhom, tu_khoa=tu_khoa, trang_thai="verified" if duyet else "cho_duyet",
              nguon=nguon or nguon_file.name, cap_nhat=str(date.today()))
    p.write_text("---\n" + yaml.safe_dump(fm, allow_unicode=True, sort_keys=False) + "---\n" + body.strip() + "\n", encoding="utf-8")
    ghi_muc_luc()
    return p


def duyet(id_: str) -> Path:
    m = lay(id_)
    if not m or not str(m["file"]).startswith(str(KHO)):
        raise SystemExit(f"{id_}: không phải mục trong knowledge/ (mục từ file nguồn mặc định đã verified)")
    fm, body = _fm(m["file"].read_text(encoding="utf-8-sig"))
    fm.update(trang_thai="verified", cap_nhat=str(date.today()))
    m["file"].write_text("---\n" + yaml.safe_dump(fm, allow_unicode=True, sort_keys=False) + "---\n" + body, encoding="utf-8")
    ghi_muc_luc()
    return m["file"]


def ghi_muc_luc() -> Path:
    """knowledge/_index.md — mỗi mục 1 dòng; nhỏ để nạp được luôn."""
    L = ["# knowledge/_index.md (tự sinh — `knowled ds`; tra bằng `knowled tim`, không nạp cả kho)", ""]
    for nh in NHOM:
        ms = [m for m in muc_kho() if m["nhom"] == nh and str(m["file"]).startswith(str(KHO))]
        if ms:
            L.append(f"## {nh}")
            L += [f"- `{m['id']}` {'✓' if m['trang_thai'] == 'verified' else '?'} {m['tieu_de']} — {m['tu_khoa']}" for m in ms]
    f = KHO / "_index.md"
    KHO.mkdir(exist_ok=True)
    f.write_text("\n".join(L) + "\n", encoding="utf-8")
    return f
