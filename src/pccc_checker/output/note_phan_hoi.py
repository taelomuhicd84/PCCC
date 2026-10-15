"""Word "Note phản hồi.docx" — tổng hợp mọi vị trí đã ghi chú trên bản vẽ, xếp theo trang, để kỹ sư kéo tới kiểm nhanh.

Nguồn: va_cham.json (va chạm, bỏ dòng `bo`) + findings.jsonl nếu có (dòng KD / XM).
Màu theo kiến trúc: Không đạt = chữ đỏ · Cần xác minh = chữ nghiêng cam · Va chạm = cam.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt, RGBColor

RED, ORANGE = RGBColor(0xC0, 0x00, 0x00), RGBColor(0xE0, 0x6C, 0x00)
TEN_FILE = "Note phản hồi.docx"


def rows_va_cham(items: list[dict]) -> list[dict]:
    """1 dòng / 1 va chạm (không lặp A≠B rồi B≠A), ghi thẳng số liệu từng nơi + mức lệch."""
    from ..pipeline.va_cham import chinh_xac, mo_ta
    out = []
    for it in items:
        if it.get("bo"):
            continue
        vts = [v.get("vi_tri") or {} for v in it["gia_tri"]]
        ok = sum(chinh_xac(vt.get("bbox")) for vt in vts)
        trangs = sorted({v.get("trang") for v in it["gia_tri"] if v.get("trang")})
        out.append({"trang": trangs[0] if trangs else None, "trangs": trangs, "ma": it["id"], "loai": "Va chạm",
                    "hm": it["hm"] or "khu đất",
                    "noi_dung": mo_ta(it) + (f". {it['ghi_chu']}" if it.get("ghi_chu") else ""), "dieu_khoan": "",
                    "vi_tri": f"khoanh {ok}/{len(vts)} vị trí trên bản vẽ" + ("" if ok == len(vts) else "; còn lại chưa chỉ được vị trí")})
    return out


def rows_findings(rows: list[dict], hm_names: dict[str, str] | None = None) -> list[dict]:
    hm_names = hm_names or {}
    out = []
    for r in rows:
        if r.get("tt") not in ("KD", "XM"):
            continue
        out.append({"trang": r.get("trang"), "ma": r.get("id"), "loai": "Không đạt" if r["tt"] == "KD" else "Cần xác minh",
                    "hm": hm_names.get(r.get("hm"), r.get("hm") or ""), "noi_dung": r.get("ly_do") or "",
                    "dieu_khoan": r.get("dieu_khoan") or r.get("rule") or "",
                    "vi_tri": "khung trên bản vẽ" if r.get("bbox") else "không có vị trí"})
    return out


def _cell(c, text: str, color=None, italic=False, bold=False, size=9):
    c.text = ""
    run = c.paragraphs[0].add_run(str(text))
    run.font.size = Pt(size)
    run.italic, run.bold = italic, bold
    if color is not None:
        run.font.color.rgb = color


def build(rows: list[dict], out_file: Path, ho_so: str = "", pdf_notes: list[str] | None = None) -> Path:
    rows = sorted(rows, key=lambda r: (r["trang"] or 10 ** 6, r["loai"], str(r["ma"])))
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.5))
    doc.add_heading("Note phản hồi", 0)
    doc.add_paragraph(f"Hồ sơ: {ho_so}\nNgày: {date.today():%d/%m/%Y}"
                      + "".join(f"\nPDF ghi chú: {p}" for p in pdf_notes or []))
    cnt = {k: sum(r["loai"] == k for r in rows) for k in ("Va chạm", "Không đạt", "Cần xác minh")}
    doc.add_paragraph(" · ".join(f"{k}: {v}" for k, v in cnt.items()))

    doc.add_heading("1. Trang cần kiểm", 1)
    pages: dict = {}
    for r in rows:
        for pg in r.get("trangs") or [r["trang"]]:
            pages.setdefault(pg, []).append(str(r["ma"]))
    pages = dict(sorted(pages.items(), key=lambda kv: kv[0] or 10 ** 6))
    p = doc.add_paragraph()
    for pg, ms in pages.items():
        run = p.add_run(f"Tr.{pg if pg else '—'}: ")
        run.bold = True
        p.add_run(", ".join(dict.fromkeys(ms)) + "\n")

    doc.add_heading("2. Chi tiết theo trang", 1)
    head = ["Trang", "Mã note", "Loại", "Hạng mục", "Nội dung", "Điều khoản", "Vị trí"]
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, head):
        _cell(c, h, bold=True)
    for r in rows:
        color = RED if r["loai"] == "Không đạt" else ORANGE
        cells = t.add_row().cells
        vals = [r["trang"] or "—", r["ma"], r["loai"], r["hm"], r["noi_dung"], r["dieu_khoan"], r["vi_tri"]]
        for i, (c, v) in enumerate(zip(cells, vals)):
            _cell(c, v, color=color if i in (2, 4) else None, italic=r["loai"] == "Cần xác minh" and i in (2, 4))
    for row in t.rows:
        for c, w in zip(row.cells, (1.4, 2.0, 2.2, 3.2, 10.5, 3.4, 2.8)):
            c.width = Cm(w)
    doc.add_paragraph("Kết quả do máy đọc bản vẽ — kỹ sư PCCC xác nhận trước khi phản hồi chính thức.").runs[0].italic = True
    out_file.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_file)
    return out_file
